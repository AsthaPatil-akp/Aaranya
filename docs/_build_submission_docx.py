from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).with_name("SUBMISSION_DRAFT.docx")


def set_run(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(27, 67, 50)
    return p


def add_p(doc, text, *, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, bold=bold, italic=italic)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)


def add_placeholder(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run(run, bold=True, color=(153, 51, 0))
    p.paragraph_format.space_after = Pt(8)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("DARUKAA.EARTH HACKATHON — FINAL SUBMISSION")
    set_run(r, bold=True, size=16, color=(27, 67, 50))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("AI Biodiversity Intelligence Chatbot")
    set_run(r, size=13, color=(27, 67, 50))

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "Draft generated from the current local codebase. Replace orange placeholders before upload. "
        "Do not paste secrets, .env contents, or API keys into this document."
    )
    set_run(r, italic=True, size=10, color=(90, 90, 90))

    add_heading(doc, "1. Project Title", 1)
    add_p(
        doc,
        "Aaranya — AI Biodiversity Intelligence Chatbot (built for the Darukaa.Earth challenge).",
    )
    add_p(
        doc,
        "The user-facing product name in the current frontend is Aaranya. The backend FastAPI title, README, "
        "and knowledge-base syntheses still use Darukaa.Earth. Both names refer to the same prototype.",
    )

    add_heading(doc, "2. Team / Candidate Information", 1)
    add_placeholder(doc, "[CANDIDATE FULL NAME]")
    add_placeholder(doc, "[EMAIL]")
    add_placeholder(doc, "[INSTITUTION / TEAM NAME]")
    add_placeholder(doc, "[ROLE — e.g. Individual contributor / Team lead]")
    add_p(doc, "Fill this section with the exact details required by the hackathon portal.")

    add_heading(doc, "3. Project Overview", 1)
    add_p(
        doc,
        "This project is a local, evidence-grounded conversational assistant for biodiversity and environmental "
        "management questions. A user describes a farm or landscape in text, optionally adds structured land "
        "details, and receives a plain-language answer that is constrained by retrieved passages.",
    )
    add_p(
        doc,
        "The system is not a generic chatbot. The pipeline extracts environmental variables, stores conversation "
        "context, retrieves internal knowledge-base chunks from ChromaDB, optionally adds OpenAlex literature, "
        "asks a local Ollama model to draft an answer, then validates and rewrites unsupported claims before "
        "showing the response.",
    )
    add_p(
        doc,
        "The current interface includes a Home page, an Intelligence Lab chat, a Knowledge Base page, and an "
        "architecture/about page. Completed Lab answers can download a client-side Biodiversity Action Plan PDF "
        "built only from the grounded response, provided land details, and used sources.",
    )

    add_heading(doc, "4. Problem Statement", 1)
    add_p(
        doc,
        "Land managers need site-specific biodiversity advice that can combine soil, climate, land use, and "
        "human-impact conditions. General-purpose language models often invent citations, timeframes, or "
        "percentages when evidence is thin. The challenge is to keep a natural conversation while refusing to "
        "present unsupported scientific claims as fact.",
    )

    add_heading(doc, "5. Solution", 1)
    add_p(
        doc,
        "The solution is a retrieval-augmented FastAPI service plus a React/Vite Lab UI. Recommendations are "
        "authored by a local LLM that is given retrieved passages, then checked by a grounding layer. Internal "
        "Darukaa syntheses and external OpenAlex records are labelled separately. If the site profile is incomplete, "
        "the system asks clarifying questions instead of inventing missing values. If evidence is insufficient, "
        "it reports that limitation rather than fabricating a paper.",
    )

    add_heading(doc, "6. Key Features", 1)
    add_p(doc, "Implemented in the current codebase:", bold=True)
    add_bullets(
        doc,
        [
            "Conversational text chat with session memory (SQLite, keyed by session_id).",
            "Optional structured input: Add Land Details form, map/geocode, and JSON object.",
            "Clarifying questions when core site variables are missing.",
            "RAG over eight seed scientific syntheses in ChromaDB, with MiniLM embeddings and BM25 hybrid rerank.",
            "Optional OpenAlex lookup when the internal knowledge base is weak or the user asks for studies.",
            "Local Ollama LLM (default llama3.2:3b); optional paid OpenAI provider exists in code but is not required.",
            "Claim-level grounding that rewrites or removes unsupported statements before display.",
            "Recommendation panel: action, why, impacted metrics, time horizon, supporting evidence, confidence.",
            "Multi-variable reasoning over soil, climate, land use, biodiversity observations, and human impact.",
            "Separate UI lists for internal knowledge-base sources and external scientific sources.",
            "Client-side Action Plan PDF from the grounded response; empty fields are omitted, not invented.",
            "Knowledge page for listing documents, searching retrieval, and admin ingest/rebuild.",
            "Health endpoint reporting embeddings, Chroma counts, and LLM availability.",
        ],
    )
    add_p(doc, "Not implemented / not claimed:", bold=True)
    add_bullets(
        doc,
        [
            "No public cloud live demo is configured in this repository.",
            "No Docker, Render, Vercel, Hugging Face Spaces, or similar deployment files are present.",
            "No Git remote is configured on the inspected working copy, so a GitHub URL is not yet available from git.",
            "Anthropic and Groq API settings exist in config.py but no corresponding LLM providers are implemented.",
            "Frontend GitHub Actions currently build the UI; they do not run npm test.",
            "The seed knowledge base is eight original syntheses, not a full copy of FAO/IPBES/IPCC reports.",
        ],
    )

    add_heading(doc, "7. Technology Stack", 1)
    add_p(doc, "Backend", bold=True)
    add_bullets(
        doc,
        [
            "Python, FastAPI, Uvicorn, Pydantic Settings",
            "SQLite (conversation memory and document catalog)",
            "ChromaDB vector store, collection darukaa_knowledge",
            "sentence-transformers all-MiniLM-L6-v2 (384-d embeddings)",
            "rank-bm25 hybrid rerank",
            "Ollama local LLM (default llama3.2:3b)",
            "Optional OpenAI chat/embeddings if LLM_PROVIDER or EMBEDDING_BACKEND is set to openai",
            "httpx for OpenAlex and optional open-access PDF text",
            "pypdf for PDF ingest",
            "pytest for backend tests",
        ],
    )
    add_p(doc, "Frontend", bold=True)
    add_bullets(
        doc,
        [
            "React 18, TypeScript, Vite",
            "react-router-dom (Home, Lab, Knowledge, About)",
            "Leaflet map for optional land-details location",
            "jsPDF for client-side action-plan download",
            "Vitest + Testing Library for frontend unit tests",
        ],
    )
    add_p(doc, "Knowledge / research APIs", bold=True)
    add_bullets(
        doc,
        [
            "Internal markdown syntheses under knowledge/sources/, catalogued in knowledge/manifest.json",
            "OpenAlex (optional; ENABLE_OPENALEX), labelled as external evidence",
            "Nominatim geocode for optional place search (GET /api/geocode)",
        ],
    )

    add_heading(doc, "8. System Architecture", 1)
    add_p(
        doc,
        "User → React Lab (POST /api/chat/stream NDJSON) or JSON POST /api/chat → environmental extraction and "
        "SQLite memory → clarification if the site picture is incomplete → context-aware retrieval query → "
        "ChromaDB + BM25 → relevance filter → optional OpenAlex → Ollama draft (json_mode=false on the stream path) → "
        "grounding validation and rewrite → filtered source lists → ChatResponse (answer, recommendation, evidence, debug).",
    )
    add_p(
        doc,
        "The Lab UI always uses the streaming endpoint. The recommendation panel and PDF are mapped from the same "
        "ChatResponse schema; there is no separate recommender service.",
    )
    add_p(
        doc,
        "If frontend/dist exists, FastAPI can also mount the built UI as static files. That is local serving only, "
        "not a hosted demo.",
    )

    add_heading(doc, "9. Knowledge Base / RAG Pipeline", 1)
    add_p(
        doc,
        "Seed files (knowledge/sources/) are ingested into SQLite document/chunk tables and embedded into ChromaDB. "
        "Chunking uses about 900 characters with 140-character overlap. Each chunk stores document name, title, topic, "
        "origin, URL/DOI when present, and whether a page number is a real PDF page.",
    )
    add_p(doc, "Seed coverage (from knowledge/manifest.json):", bold=True)
    add_bullets(
        doc,
        [
            "Soil health: soil organic carbon, pH, moisture, below-ground biodiversity; cover crops and residue in semi-arid farming.",
            "Climate: rainfall/drought and species survival; temperature, drought, and vegetation stress.",
            "Land use / land cover: monoculture vs agroforestry and habitat complexity; deforestation and fragmentation; land-use change across grasslands, forests, and agriculture.",
            "Human impact: pollution, urban expansion, and species richness.",
            "Biodiversity indicators appear across these syntheses (pollinators, habitat complexity, species richness, species survival).",
        ],
    )
    add_p(
        doc,
        "Retrieval hybrid score is approximately 0.70 dense similarity + 0.20 BM25 + 0.10 token overlap. Weak hits "
        "below RELEVANCE_THRESHOLD (default 0.42) are not treated as evidence. Internal and external passages are "
        "kept in separate lists. Unused off-topic retrieved documents (for example unused deforestation chunks) are "
        "filtered from the user-facing sources.",
    )
    add_p(
        doc,
        "These seed files are Darukaa syntheses that point at public institutional science (FAO, IPBES, IPCC, USDA NRCS, CBD). "
        "They are not the original institutional reports.",
    )

    add_heading(doc, "10. Database / Schema", 1)
    add_p(doc, "SQLite file: data/darukaa.sqlite (runtime; gitignored).", bold=True)
    add_bullets(
        doc,
        [
            "conversations(id, created_at, updated_at, context_json) — EnvironmentalContext JSON per session.",
            "messages(id, session_id, role, content, created_at) — chat history, not mixed across sessions.",
            "documents(...) — catalog: name, source, type, topic, pages, checksum, title, authors, year, doi, url, origin.",
            "chunks(...) — text chunks with page, chunk_index, topic, page_is_real.",
        ],
    )
    add_p(doc, "ChromaDB: data/chroma, collection darukaa_knowledge. One vector per chunk (384-d MiniLM).", bold=True)
    add_p(
        doc,
        "Chat JSON schema includes EnvironmentalContext (location, soil, land, biodiversity, climate, human_impact), "
        "RecommendationBlock (action, why_it_works, environmental_relationships, impacted_metrics, time_horizon, "
        "uncertainty, confidence, items, supporting_evidence), EvidenceItem, and knowledge_status.",
    )

    add_heading(doc, "11. AI / Reasoning Workflow", 1)
    add_bullets(
        doc,
        [
            "Extract variables from free text and optional structured fields; merge into session context.",
            "If fewer than about three core variables are known for an advisory question, ask targeted clarifying questions.",
            "Build a retrieval query from known variables plus the user message.",
            "Retrieve and filter internal passages; call OpenAlex only when configured and needed.",
            "Pass INTERNAL KNOWLEDGE BASE and EXTERNAL SCIENTIFIC EVIDENCE separately into the LLM prompt, with optional candidate interventions (not treated as facts).",
            "Draft an answer. On /api/chat/stream this is conversational text; on /api/chat it is structured JSON that is then composed into the user-facing answer.",
            "Grounding: verify claims against retrieved passages; rewrite or omit unsupported statements (including invented crop rotation, exact year spans, or pesticide claims when pesticide_use is none).",
            "Complete the recommendation schema from the grounded answer, matched candidates, and used source titles.",
            "Return assistant_message plus recommendation, metrics, time horizon, confidence, and source lists.",
        ],
    )
    add_p(
        doc,
        "Heuristic soil/habitat scores in the Lab sidebar are labelled as heuristic readings of the inputs, not validated predictions.",
    )

    add_heading(doc, "12. Input and Output", 1)
    add_p(doc, "Input", bold=True)
    add_bullets(
        doc,
        [
            "Mandatory: free-text message in the Lab composer (POST /api/chat or /api/chat/stream).",
            "Optional structured: structured object or structured_json; Add Land Details (farm size, location, lat/lng, crop, land use, soil pH, SOC, moisture, rainfall, temperature, pesticide_use, biodiversity observations).",
            "Optional debug=true to include retrieval traces (used by the Lab developer panel).",
        ],
    )
    add_p(doc, "Output", bold=True)
    add_bullets(
        doc,
        [
            "assistant_message: grounded conversational answer.",
            "clarifying_questions when mode is clarification.",
            "recommendation: what to do, why, environmental relationships, impacted metrics (often “potentially affected”), time horizon (or an explicit “no specific timeframe is supported” statement), confidence, supporting evidence titles.",
            "kb_evidence and external_evidence kept apart, with evidence_level (for example abstract vs full_text).",
            "knowledge_status: grounded_in_knowledge_base | grounded_in_external_evidence | grounded_in_kb_and_external | insufficient_evidence | awaiting_clarification.",
            "Optional Action Plan PDF with site profile, assessment, recommended actions (checkboxes), monitoring checklist, sources, and limitations.",
        ],
    )

    add_heading(doc, "13. Scientific Evidence / Sources", 1)
    add_p(
        doc,
        "User-facing answers cite retrieved document titles. Internal syntheses are labelled as internal knowledge-base "
        "documents, not as original FAO/IPBES papers. External OpenAlex items, when used, include title, year, DOI/URL "
        "when present, and evidence level. The PDF and Lab source panels list only sources used in the grounded answer.",
    )
    add_p(
        doc,
        "Related institutional URLs stored in the seed manifest include FAO soils/land portals, IPBES, USDA NRCS soil health, and CBD. Those URLs are metadata for the syntheses; the chatbot does not scrape those sites at runtime except via OpenAlex when enabled.",
    )

    add_heading(doc, "14. How the Project Meets Hackathon Requirements", 1)
    add_p(doc, "Conversational biodiversity intelligence: implemented in the Lab with FastAPI chat/stream.")
    add_p(doc, "Retrievable knowledge layer: implemented (MiniLM + ChromaDB + BM25 + SQLite catalog).")
    add_p(doc, "Coverage of soil, land use, biodiversity, climate, human impact: implemented via eight seed syntheses and EnvironmentalContext fields.")
    add_p(doc, "Conversational context and clarifying questions: implemented (session memory, history_window default 6, clarification mode).")
    add_p(doc, "Recommendations with what / why / metric / source: implemented in RecommendationBlock, Lab panel, and PDF.")
    add_p(doc, "Multi-variable reasoning: implemented in prompting and environmental_relationships; farm-profile tests check interacting soil/climate/land-use conditions.")
    add_p(doc, "Text input mandatory, structured input supported: implemented.")
    add_p(doc, "Output includes recommendation, impacted metrics, time horizon, confidence: implemented. Time horizons are omitted or marked unsupported unless evidence supports them.")
    add_p(
        doc,
        "Evaluation themes: grounding tests rewrite unsupported claims; retrieval tests check relevance filtering; Lab UI separates sources and shows uncertainty copy when fields are empty.",
    )

    add_heading(doc, "15. Local Setup", 1)
    add_p(doc, "Requirements: Python 3.11+ recommended, Node.js 20 for the UI, Ollama with llama3.2:3b for live chat.")
    add_p(doc, "Backend", bold=True)
    add_bullets(
        doc,
        [
            "python -m venv .venv && .venv\\Scripts\\activate",
            "pip install -r backend/requirements.txt",
            "copy .env.example .env  (then set ADMIN_API_TOKEN; keep secrets local)",
            "Install Ollama, run ollama serve, ollama pull llama3.2:3b",
            "From backend: uvicorn app.main:app --reload --port 8000",
            "Health: http://127.0.0.1:8000/api/health — confirm llm_available=true only when Ollama is reachable",
        ],
    )
    add_p(doc, "Frontend", bold=True)
    add_bullets(
        doc,
        [
            "cd frontend && npm install && npm run dev",
            "Open http://127.0.0.1:5173 (Vite proxies /api to port 8000). The app is configured to land on Home first.",
        ],
    )
    add_p(doc, "Tests", bold=True)
    add_bullets(
        doc,
        [
            "From repo root: pytest -q  (mocks Ollama; does not need a live model)",
            "cd frontend && npm test",
            "Optional live model: RUN_OLLAMA_INTEGRATION=1 pytest -q tests/test_ollama.py::test_live_ollama_optional",
        ],
    )

    add_heading(doc, "16. Deployment", 1)
    add_p(
        doc,
        "Not deployed. This repository has no Dockerfile, docker-compose, or cloud hosting config. Reviewers should run the local setup. If a live URL is later published, replace the placeholder below. Do not claim a hosted demo until it exists.",
    )
    add_placeholder(doc, "[LIVE DEMO URL] — currently: not deployed")

    add_heading(doc, "17. CI/CD", 1)
    add_p(
        doc,
        "GitHub Actions workflow .github/workflows/ci.yml runs on push/PR to main or master:",
    )
    add_bullets(
        doc,
        [
            "Backend job: Python 3.11, pip install backend/requirements.txt, pytest -q, with OpenAlex disabled and a test admin token.",
            "Frontend job: Node 20, npm install, npm run build. It does not currently run npm test.",
        ],
    )
    add_p(
        doc,
        "There is no CD/deploy job. CI will only run after the project is pushed to GitHub.",
    )

    add_heading(doc, "18. GitHub Repository", 1)
    add_placeholder(doc, "[GITHUB REPOSITORY URL]")
    add_p(
        doc,
        "Inspected working copy: branch main, latest committed snapshot “Ground Lab chat in local RAG and add optional land details.” "
        "No git remote was configured at inspection time. Substantial later work (grounding rewrite, land-details geocode/map, "
        "recommendation-panel mapping, Aaranya branding, action-plan PDF) exists in the local working tree and must be committed "
        "and pushed before reviewers can see it on GitHub.",
    )

    add_heading(doc, "19. Live Demo", 1)
    add_placeholder(doc, "[LIVE DEMO URL] — not applicable unless you host the app")
    add_p(
        doc,
        "Local demo for judges: start Ollama + API + Vite, then open the Home page and Intelligence Lab. Suggested script is in README.md: incomplete farm question → clarifying questions → farm profile → inspect sources, metrics, time horizon, confidence → follow-up using memory → out-of-corpus question should not invent papers → download Action Plan PDF from a completed answer.",
    )

    add_heading(doc, "20. Demo Credentials / Reviewer Notes", 1)
    add_p(
        doc,
        "Chat and retrieval do not require a user login. Knowledge ingest, rebuild, and GET /api/conversation/{id} require header X-Admin-Token matching the server ADMIN_API_TOKEN.",
    )
    add_placeholder(doc, "[DEMO ADMIN TOKEN IF YOU HOST A REVIEW INSTANCE — use a throwaway token, never a personal API key]")
    add_p(
        doc,
        "Do not put production OpenAI keys, private Ollama tunnels, or real .env values in this document. If reviewers only run locally, they copy .env.example and set their own ADMIN_API_TOKEN.",
    )
    add_p(doc, "Reviewer notes:", bold=True)
    add_bullets(
        doc,
        [
            "Live chat needs a running Ollama model. Automated pytest does not.",
            "OpenAlex needs internet when ENABLE_OPENALEX=true. Internal KB answers can run without it.",
            "Empty land-detail fields stay unknown; the model is instructed not to invent them.",
            "UI product name is Aaranya; documents and API title still say Darukaa.Earth.",
        ],
    )

    add_heading(doc, "21. Known Limitations", 1)
    add_bullets(
        doc,
        [
            "Knowledge base is a small set of eight educational syntheses, not a comprehensive literature corpus.",
            "External evidence is often abstract-level unless an open-access PDF is fetched.",
            "Local 3B-class models can still drift; grounding reduces but does not eliminate residual error.",
            "No hosted demo, so judges must run locally unless you deploy before the deadline.",
            "GitHub remote and latest local features were not fully published at inspection time.",
            "Time horizons and metric directions are conservative (“potentially affected”) when evidence is weak.",
            "Geocoding depends on Nominatim availability.",
            "Admin token in the Knowledge UI is stored in browser sessionStorage; treat demo tokens as disposable.",
        ],
    )

    add_heading(doc, "22. Future Improvements", 1)
    add_bullets(
        doc,
        [
            "Publish the repository and a read-only live demo with a disposable admin token.",
            "Commit remaining local features so GitHub matches the demo.",
            "Add frontend npm test to CI; add a simple CD pipeline if hosting.",
            "Expand the knowledge base with additional licensed or original syntheses and more full-text papers.",
            "Align backend APP_NAME / README branding with the Aaranya UI name, or revert the UI name if the challenge requires Darukaa.Earth everywhere.",
            "Optional: Docker Compose for one-command Ollama + API + UI.",
        ],
    )

    add_heading(doc, "A. Final checklist before submission", 1)
    add_bullets(
        doc,
        [
            "Commit and push all intended features to GitHub; confirm the public repo matches the demo.",
            "Replace [GITHUB REPOSITORY URL] with the real HTTPS URL.",
            "If hosted, replace [LIVE DEMO URL]; otherwise state clearly that the demo is local-only.",
            "Fill candidate/team fields required by the portal.",
            "Run pytest -q and frontend npm test on a clean checkout.",
            "Run the judge demo script once with Ollama actually available.",
            "Confirm .env is not in git (it is gitignored). Confirm no API keys are in frontend source.",
            "If sharing Knowledge ingest, give a throwaway ADMIN_API_TOKEN, not a personal secret.",
            "Decide whether to keep the Aaranya UI name or standardise on Darukaa.Earth for judging.",
        ],
    )

    add_heading(doc, "B. Things still needed from you", 1)
    add_bullets(
        doc,
        [
            "[GITHUB REPOSITORY URL] — create/push origin; none was configured on the inspected clone.",
            "[LIVE DEMO URL] — optional; not currently deployed.",
            "[CANDIDATE FULL NAME], [EMAIL], [INSTITUTION / TEAM NAME].",
            "[DEMO ADMIN TOKEN] only if you host a public Knowledge ingest endpoint.",
            "Commit unpublished local work (PDF export, grounding, land-details map, recommendation mapping, Aaranya branding) if those are part of the judged demo.",
        ],
    )

    add_heading(doc, "C. Things you should NOT upload to GitHub", 1)
    add_bullets(
        doc,
        [
            ".env, .env.local, credentials*.json, *.pem, *.key",
            "data/ (SQLite, Chroma runtime files)",
            ".venv/, node_modules/, frontend/dist/",
            "Real OPENAI_API_KEY, ANTHROPIC_API_KEY, GROQ_API_KEY, or production ADMIN_API_TOKEN",
            "Private conversation databases or personal land coordinates you do not want public",
        ],
    )

    add_heading(doc, "D. Things you should NOT put in the Word document", 1)
    add_bullets(
        doc,
        [
            "Any value copied from a live .env file",
            "OpenAI/Anthropic/Groq API keys",
            "Personal passwords or Ollama cloud keys",
            "Internal admin tokens you still use privately",
            "Invented live URLs, datasets, or model accuracy percentages",
        ],
    )

    add_heading(doc, "E. Missing or at-risk hackathon requirements", 1)
    add_p(
        doc,
        "Functional requirements for RAG, conversation, structured+text input, multi-variable reasoning, and recommendation fields are implemented in code. Submission-process gaps are the main risk:",
    )
    add_bullets(
        doc,
        [
            "No GitHub URL until you push the repo — evaluators cannot review code without it.",
            "No live demo URL — acceptable if the brief allows local setup, but then setup steps must be flawless.",
            "Uncommitted local features will not be visible to judges if you only push the current last commit.",
            "CI frontend does not run unit tests; backend CI exists but only after GitHub is connected.",
            "Do not describe seed syntheses as original FAO/IPBES publications.",
            "Optional OpenAlex is not a substitute for the internal knowledge layer; the internal layer is the primary RAG corpus.",
        ],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
